"""A small, honest accuracy check. Run with: python eval.py

This is not a toy — every expected value below was verified against the
actual sample data (sample_docs/ + both tables in sample_data/), not
guessed. If you swap in your own documents/spreadsheets, write your own
cases the same way: compute the right answer yourself first, then check
the agent reaches it.
"""
from __future__ import annotations

import io
import os
import tempfile

from agent import RAGAgent

CASES = [
    # --- text RAG: checked by which source file(s) come back ---
    {
        "question": "What is safety stock?",
        "check": lambda r: "safety_stock.txt" in r["sources"],
        "description": "text RAG retrieves the right source document",
    },
    {
        "question": "What does DDP mean?",
        "check": lambda r: "incoterms.txt" in r["sources"],
        "description": "text RAG retrieves the right source document (2)",
    },
    {
        "question": "What is stranded inventory?",
        "check": lambda r: "fba_inventory_health.txt" in r["sources"],
        "description": "text RAG retrieves the right source document (3)",
    },
    {
        "question": "zzxxqqnonexistentgarbagewordxyz",
        "check": lambda r: "No matching passages" in r["answer"],
        "description": "text RAG honestly reports no match instead of guessing",
    },
    # --- e-commerce business-model docs: commercial side, catalog, PnL,
    # contracts, and the cross-document links between them ---
    {
        "question": "Explain the 3P model",
        "check": lambda r: "ecommerce_business_models.txt" in r["sources"],
        "description": "retrieves the 3P/1P/DTC explainer even though a less-relevant "
        "chunk elsewhere uses the generic word 'model' more densely (a real TF-IDF "
        "ranking quirk — this checks the right doc is still SURFACED, not ranked #1)",
    },
    {
        "question": "What causes stranded inventory and how does it affect forecasting?",
        "check": lambda r: "catalog_management.txt" in r["sources"]
        and "catalog_supply_chain_connection.txt" in r["sources"],
        "description": "graph-links the catalog root cause with the forecasting consequence across two files",
    },
    {
        "question": "What is contribution margin?",
        "check": lambda r: "ecommerce_pnl.txt" in r["sources"],
        "description": "retrieves the P&L document for a margin/economics question",
    },
    {
        "question": "What happens during brand onboarding?",
        "check": lambda r: "brand_contracts_onboarding.txt" in r["sources"],
        "description": "retrieves the onboarding process document",
    },
    # --- structured: exact lookup ---
    {
        "question": "Tell me about WPB-1004",
        "check": lambda r: "BLR2" in r["answer"] and "WPB-1004" in r["answer"],
        "description": "exact SKU lookup returns the right row",
    },
    {
        "question": "What is B0CXP7Q8R9",
        "check": lambda r: "Stranded" in r["answer"],
        "description": "exact ASIN lookup returns the right row",
    },
    # --- structured: single-condition filters ---
    {
        "question": "Which SKUs have Quantity below 30",
        "check": lambda r: all(s in r["answer"] for s in ["WPB-1002", "WPB-1004", "WPB-1006"])
        and "WPB-1001" not in r["answer"],
        "description": "numeric filter (<) returns exactly the right rows",
    },
    {
        "question": "How many SKUs have Status Low",
        "check": lambda r: "3 matching row" in r["answer"],
        "description": "categorical filter + count is exact",
    },
    # --- structured: AND / OR ---
    {
        "question": "Which SKUs have Quantity below 30 and FC is BLR2",
        "check": lambda r: "WPB-1004" in r["answer"]
        and "WPB-1002" not in r["answer"]
        and "WPB-1006" not in r["answer"],
        "description": "AND across two conditions narrows correctly (the regression this harness exists to catch)",
    },
    {
        "question": "Which SKUs have Status Low or Status Stranded",
        "check": lambda r: all(s in r["answer"] for s in ["WPB-1002", "WPB-1004", "WPB-1006", "WPB-1008"])
        and "WPB-1001" not in r["answer"],
        "description": "OR across two conditions widens correctly",
    },
    # --- structured: aggregations ---
    {
        "question": "What is the total quantity",
        "check": lambda r: "479" in r["answer"],
        "description": "sum aggregation over the whole table",
    },
    {
        "question": "What is the average ReorderPoint",
        "check": lambda r: "51.25" in r["answer"],
        "description": "mean aggregation over the whole table",
    },
    {
        "question": "What is the total quantity by FC",
        "check": lambda r: all(s in r["answer"] for s in ["BLR2", "133", "BOM1", "136", "DEL3", "210"]),
        "description": "grouped sum aggregation",
    },
    {
        "question": "What is the total quantity of SKUs with Status Low",
        "check": lambda r: "56" in r["answer"],
        "description": "aggregation combined with a filter (the other regression this harness exists to catch)",
    },
    # --- multi-table routing: a second, larger structured table ingested
    # alongside the first one (fba_inventory_detailed.xlsx, 90 rows, 27
    # columns) — these exist specifically to catch the aggregation engine
    # picking the WRONG table just because a column name happens to
    # substring-match the question (e.g. "SKU" matching inside "SKUs").
    {
        "question": "What is the total Inventory Value",
        "check": lambda r: "177209.09" in r["answer"] and r["sources"] == ["fba_inventory_detailed.xlsx"],
        "description": "aggregation picks the correct table out of several ingested tables",
    },
    {
        "question": "How many SKUs have Stockout Risk HIGH",
        "check": lambda r: "12 matching row" in r["answer"],
        "description": "categorical filter + count on the second table",
    },
    {
        "question": "What is the total Inventory Value of SKUs with Stockout Risk HIGH",
        "check": lambda r: "8692.96" in r["answer"] and r["sources"] == ["fba_inventory_detailed.xlsx"],
        "description": "aggregation + filter doesn't get hijacked by a same-named column in a different table (the multi-table regression this harness exists to catch)",
    },
    # --- synonym/abbreviation matching (synonyms.py) ---
    {
        "question": "Which SKUs have units below 30",
        "check": lambda r: "WPB-1002" in r["answer"] and "WPB-1004" in r["answer"] and "WPB-1006" in r["answer"],
        "description": "synonym match connects 'units' to the 'Quantity' column, not just the literal name",
    },
    # --- column-vs-column comparison (no literal number at all) ---
    {
        "question": "How many SKUs have Available below Reorder Point",
        "check": lambda r: "47 matching row" in r["answer"],
        "description": "column-vs-column comparison (Available < Reorder Point), not against a fixed number",
    },
]


def run_core_cases() -> tuple[int, int]:
    agent = RAGAgent()
    agent.load()

    passed = 0
    for case in CASES:
        result = agent.query(case["question"])
        ok = case["check"](result)
        status = "PASS" if ok else "FAIL"
        passed += ok
        print(f"[{status}] {case['description']}")
        if not ok:
            print(f"         Q: {case['question']}")
            print(f"         A: {result['answer'][:200]}")
    return passed, len(CASES)


def run_upload_case() -> tuple[int, int]:
    """Exercises the actual Flask /api/upload-data endpoint — this is the
    one behavior the core-engine cases above can't cover, since it's about
    app.py's upload-replaces-demo-data behavior, not structured_query.py.
    Built this way (a real CSV in memory, posted through the real test
    client) specifically to catch the regression that motivated it: a live
    upload getting silently answered from leftover demo data instead of the
    file that was just uploaded.
    """
    import app as app_module

    client = app_module.app.test_client()

    csv_bytes = b"SKU,Qty Avail,Reorder Lvl\nX1,45,20\nX2,8,20\nX3,120,50\nX4,3,10\n"

    resp = client.post(
        "/api/upload-data",
        data={"file": (io.BytesIO(csv_bytes), "eval_upload_test.csv")},
        content_type="multipart/form-data",
    )
    upload_ok = resp.status_code == 200 and "eval_upload_test.csv" in [
        t["name"] for t in resp.get_json().get("tables", [])
    ]
    print(f"[{'PASS' if upload_ok else 'FAIL'}] live upload endpoint accepts a new CSV")

    resp = client.post("/api/ask", json={"question": "Which SKUs have stock on hand below 10"})
    answer_data = resp.get_json()
    replace_ok = (
        answer_data.get("sources") == ["eval_upload_test.csv"]
        and "X2" in answer_data.get("answer", "")
        and "X4" in answer_data.get("answer", "")
    )
    print(
        f"[{'PASS' if replace_ok else 'FAIL'}] uploaded file answers via synonym match "
        f"('stock on hand' -> 'Qty Avail'), replacing demo data rather than merging with it "
        f"(the regression this check exists to catch)"
    )

    # The real endpoint calls agent.structured.save(), which just overwrote
    # data/structured_tables.pkl with ONLY the test upload — restore the
    # actual demo tables so this stays safe to run repeatedly and doesn't
    # break the core cases (which load from that same file) on the next run.
    sample_data_dir = os.path.join(os.path.dirname(__file__), "sample_data")
    app_module.agent.structured.tables = {}
    app_module.agent.structured.ingest_folder(sample_data_dir)
    app_module.agent.structured.save()

    return int(upload_ok) + int(replace_ok), 2


def run_docx_and_chart_cases() -> tuple[int, int]:
    """Covers two features added together: live .docx upload (additive,
    not replacing — the opposite behavior from the spreadsheet upload on
    purpose, see app.py's upload_docs docstring) and chart generation for
    grouped aggregations.
    """
    import docx

    import app as app_module

    client = app_module.app.test_client()
    passed = 0

    doc = docx.Document()
    doc.add_paragraph("Eval Return Policy")
    doc.add_paragraph("Returns must be initiated within 30 days of delivery for a full refund.")
    docx_path = os.path.join(tempfile.gettempdir(), "eval_test_policy.docx")
    doc.save(docx_path)

    with open(docx_path, "rb") as f:
        resp = client.post(
            "/api/upload-docs",
            data={"file": (f, "eval_test_policy.docx")},
            content_type="multipart/form-data",
        )
    upload_ok = resp.status_code == 200 and resp.get_json().get("filename") == "eval_test_policy.docx"
    print(f"[{'PASS' if upload_ok else 'FAIL'}] live .docx upload is read and ingested")
    passed += upload_ok

    resp = client.post("/api/ask", json={"question": "What is the return window for a full refund?"})
    answer_data = resp.get_json()
    docx_answer_ok = "eval_test_policy.docx" in answer_data.get("sources", []) and "30 days" in answer_data.get(
        "answer", ""
    )
    print(f"[{'PASS' if docx_answer_ok else 'FAIL'}] uploaded .docx content is actually retrievable")
    passed += docx_answer_ok

    resp = client.post("/api/ask", json={"question": "What is safety stock?"})
    additive_ok = "safety_stock.txt" in resp.get_json().get("sources", [])
    print(
        f"[{'PASS' if additive_ok else 'FAIL'}] document upload is additive — existing "
        "text corpus still answers (the opposite of the spreadsheet upload's replace behavior, on purpose)"
    )
    passed += additive_ok

    resp = client.post("/api/ask", json={"question": "What is the total quantity by FC"})
    chart_data = resp.get_json()
    chart_ok = bool(chart_data.get("chart")) and chart_data["chart"].startswith("data:image/png;base64,")
    print(f"[{'PASS' if chart_ok else 'FAIL'}] grouped aggregation returns an actual chart image, not just text")
    passed += chart_ok

    resp = client.post("/api/ask", json={"question": "What is the total quantity"})
    no_chart_data = resp.get_json()
    no_chart_ok = no_chart_data.get("chart") is None
    print(
        f"[{'PASS' if no_chart_ok else 'FAIL'}] a single (non-grouped) number correctly has no chart — "
        "nothing to plot"
    )
    passed += no_chart_ok

    return passed, 5


def main():
    core_passed, core_total = run_core_cases()
    upload_passed, upload_total = run_upload_case()
    docx_passed, docx_total = run_docx_and_chart_cases()

    passed = core_passed + upload_passed + docx_passed
    total = core_total + upload_total + docx_total
    print(f"\n{passed}/{total} passed")


if __name__ == "__main__":
    main()
